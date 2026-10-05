"""
core/docx_export.py — 会议纪要 Markdown → docx 导出（QA-R1 DEF-02 / R6 AC-1）

产品裁决（QA-R1 缺陷派工 §0.3）：保留 docx 交付格式，AC-1 口径不收敛为
md-only；python-docx 已在依赖（requirements.txt:38），仅用于知识库读取的历史保持不变。

映射口径（Done 判据：Word/WPS 可正常打开、结构完整、中文不乱码）：
  - `#{1,6}` 标题        → Word 标题 1–6（Heading styles）
  - `**粗体**` / `*斜体*` → run 级 bold / italic（段内混排）
  - `- ` / `1. ` 列表     → List Bullet / List Number 样式
  - `【说话人N】…` 段落   → 普通段落（含说话人前缀原文，时间戳如存在一并保留）
  - `---` 分隔线 / 空行   → 段间距（不产生垃圾段落）
  中文字体：Normal 与标题样式的 w:eastAsia 显式设为宋体，避免 Word 回退渲染异常。
"""

import logging
import re
from pathlib import Path

logger = logging.getLogger(__name__)

# Markdown 水平线（整行 --- / *** / ___）
_HR_RE = re.compile(r"^\s*([-*_])\s*(?:\1\s*){2,}$")
_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_BULLET_RE = re.compile(r"^\s*[-*]\s+(.*)$")
_ORDERED_RE = re.compile(r"^\s*(\d+)[.、]\s+(.*)$")
# 段内混排：**bold** 或 *italic*（非贪婪；bold 优先）
_INLINE_RE = re.compile(r"(\*\*.+?\*\*|\*[^*\n]+?\*)")


def _apply_cjk_fonts(doc) -> None:
    """为 Normal 与 Heading 1–6 / List 样式设置 eastAsia 中文字体（宋体），防乱码。"""
    from docx.oxml.ns import qn

    style_names = ["Normal"] + [f"Heading {i}" for i in range(1, 7)] + ["List Bullet", "List Number"]
    for name in style_names:
        try:
            style = doc.styles[name]
        except KeyError:
            continue
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.get_or_add_rFonts()
        rfonts.set(qn("w:eastAsia"), "宋体")


def _add_runs_with_inline(paragraph, text: str) -> None:
    """把含 **bold** / *italic* 标记的文本拆成带格式的 runs（其余按原文纯文本）。"""
    for token in _INLINE_RE.split(text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**") and len(token) > 4:
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("*") and token.endswith("*") and len(token) > 2:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        else:
            paragraph.add_run(token)


def markdown_to_docx(md_text: str, out_path: str | Path) -> str:
    """把纪要 Markdown 文本渲染为 docx 文件，返回落盘路径。

    逐行解析（纪要模板为行式结构，无表格/代码块/嵌套列表需求）；
    不认识的语法一律按段落原文输出——保证「结构完整」优先于「语法完备」。
    """
    from docx import Document

    out_path = Path(out_path)
    doc = Document()
    _apply_cjk_fonts(doc)

    for raw_line in md_text.splitlines():
        line = raw_line.rstrip()
        stripped = line.strip()
        if not stripped:
            continue
        if _HR_RE.match(stripped):
            continue

        m = _HEADING_RE.match(line)
        if m:
            level = min(len(m.group(1)), 6)
            para = doc.add_heading("", level=level)
            _add_runs_with_inline(para, m.group(2).strip())
            continue

        m = _ORDERED_RE.match(line)
        if m:
            para = doc.add_paragraph(style="List Number")
            _add_runs_with_inline(para, m.group(2).strip())
            continue

        m = _BULLET_RE.match(line)
        if m:
            para = doc.add_paragraph(style="List Bullet")
            _add_runs_with_inline(para, m.group(1).strip())
            continue

        para = doc.add_paragraph()
        _add_runs_with_inline(para, stripped)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    logger.info(f"[导出] 纪要 docx 已生成: {out_path}")
    return str(out_path)


def export_task_docx(task: dict, out_path: str | Path) -> str:
    """按任务记录导出 docx：优先读取已落盘纪要 md，回退任务 summary 字段。

    抛 FileNotFoundError（调用方转 404）当纪要尚不存在。
    """
    output_path = task.get("output_path")
    if output_path and Path(output_path).exists():
        md_text = Path(output_path).read_text(encoding="utf-8")
    else:
        summary = task.get("summary")
        if not summary:
            raise FileNotFoundError("Minutes file not found")
        md_text = summary
    return markdown_to_docx(md_text, out_path)
