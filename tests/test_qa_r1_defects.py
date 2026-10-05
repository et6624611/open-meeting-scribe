"""
tests/test_qa_r1_defects.py — QA-R1 缺陷修复回归（DEF-QA-R1-01/02/03/04 后端侧）

依据 QA-R1 缺陷派工 §2。主 venv 无 funasr/torch，全程打桩
（打桩约定沿用 tests/test_local_engine_r1r2r3.py）；真实离线加载/推理由
macOS 活体探针 tests/eval_local_engine/qa_r1fix_offline_probe.py 补证（Spike venv）。

覆盖 AC：
- AC-D1: 模型加载入参为本地文件系统绝对路径（不再直传 ModelScope ID）；权重缺失时
         报可行动 models_missing 错误，绝不隐式回连网络（DEF-01 Done 判据 1/2/3）
- AC-D3: worker 协议新增 warmup 命令，预热加载 ASR/声纹单例；宿主启动健康后自动
         后台预热；留量声明覆盖冻结态冷 import（DEF-03 方案①）
- AC-D4: retranscribe 支持显式 engine_override=cloud 并透传至 transcribe_audio；
         拒绝白名单外取值；云端未配置时报可行动 400；笔记/音频状态不重置
         （DEF-04 后端侧）
- AC-D2: md→docx 结构映射（标题/列表/粗斜体/中文）与下载端点 format 参数
         （DEF-02 后端侧）
"""

import sys
import types
import wave
from pathlib import Path

import pytest

import core.engine_worker as w
from core.model_registry import MODEL_BY_KEY

# ============================================================
# 公共桩
# ============================================================


class _FakeAutoModel:
    """记录 AutoModel 入参的替身（模拟本地路径加载，零网络）。"""

    instances = []

    def __init__(self, **kwargs):
        self.kwargs = kwargs
        _FakeAutoModel.instances.append(self)

    def generate(self, **kwargs):
        return [{"text": "测试文本", "sentence_info": []}]


@pytest.fixture()
def fake_funasr(monkeypatch):
    """向 sys.modules 注入 funasr 替身，捕获 AutoModel 全部入参。"""
    mod = types.ModuleType("funasr")
    mod.AutoModel = _FakeAutoModel
    monkeypatch.setitem(sys.modules, "funasr", mod)
    _FakeAutoModel.instances = []
    monkeypatch.setattr(w, "_ASR_MODEL", None)
    monkeypatch.setattr(w, "_SPK_MODEL", None)
    return _FakeAutoModel


@pytest.fixture()
def fake_model_cache(monkeypatch, tmp_path):
    """伪造完整本地权重目录布局：{cache}/hub/{modelscope_id}/config.yaml。

    MODELSCOPE_CACHE 绝对路径锚点即 _local_model_path 的运行时真源
    （engine_host.build_engine_env 已将其 resolve 为绝对路径）。
    """
    cache = tmp_path / "models"
    for meta in MODEL_BY_KEY.values():
        root = cache / "hub" / meta["modelscope_id"]
        root.mkdir(parents=True)
        (root / "config.yaml").write_text("model: stub\n", encoding="utf-8")
    monkeypatch.setenv("MODELSCOPE_CACHE", str(cache))
    # get_model_dir 读 settings（可能为相对默认值）——锚点分支由 MODELSCOPE_CACHE 承接
    import core.model_registry as mr

    monkeypatch.setattr(mr, "get_model_dir", lambda: Path("data/models"))
    return cache


def _write_wav(path, seconds=1.0, sample_rate=16000):
    n = int(seconds * sample_rate)
    with wave.open(str(path), "wb") as wf:
        wf.setnchannels(1)
        wf.setsampwidth(2)
        wf.setframerate(sample_rate)
        wf.writeframes(b"\x00\x00" * n)
    return str(path)


# ============================================================
# AC-D1 — DEF-QA-R1-01：本地路径解析、禁隐式出网、可行动缺失错误
# ============================================================


class TestLocalModelPath:
    def test_path_is_absolute_under_registry_layout(self, fake_model_cache):
        p = Path(w._local_model_path("seaco-paraformer-zh"))
        assert p.is_absolute()
        # 布局与 model_registry 同源：{cache}/hub/{modelscope_id}
        assert p == fake_model_cache / "hub" / MODEL_BY_KEY["seaco-paraformer-zh"]["modelscope_id"]

    def test_missing_weights_raises_actionable_error(self, monkeypatch, tmp_path):
        monkeypatch.setenv("MODELSCOPE_CACHE", str(tmp_path / "empty-cache"))
        import core.model_registry as mr

        monkeypatch.setattr(mr, "get_model_dir", lambda: Path("data/models"))
        with pytest.raises(w.ModelsNotReadyError) as ei:
            w._local_model_path("cam-plus")
        msg = str(ei.value)
        # 可行动：指明模型 key、预期目录、引导下载
        assert "cam-plus" in msg and "下载" in msg and str(tmp_path) in msg

    def test_asr_model_loads_with_fs_paths_not_ids(self, fake_funasr, fake_model_cache):
        w._get_asr_model()
        assert len(fake_funasr.instances) == 1
        kwargs = fake_funasr.instances[0].kwargs
        for role in ("model", "vad_model", "spk_model", "punc_model"):
            value = kwargs[role]
            assert Path(value).is_absolute(), f"{role} 必须是绝对路径（禁 ModelScope ID 直传）"
            assert Path(value).is_dir()
            # 不是 registry ID 原文（ID 含 '/' 但并非存在目录的绝对路径）
            assert not value.startswith("iic/")
        # 观测口径不变：_oms_model_ids 仍记录 ModelScope ID（人类可读日志）
        inst = fake_funasr.instances[0]
        assert inst._oms_model_ids["model"] == MODEL_BY_KEY["seaco-paraformer-zh"]["modelscope_id"]

    def test_ct_punc_optional_skipped_when_absent(self, fake_funasr, fake_model_cache, monkeypatch):
        import shutil

        shutil.rmtree(fake_model_cache / "hub" / MODEL_BY_KEY["ct-punc"]["modelscope_id"])
        w._get_asr_model()
        kwargs = fake_funasr.instances[0].kwargs
        assert "punc_model" not in kwargs
        assert kwargs["model"]  # 必需件仍完整

    def test_spk_model_loads_with_fs_path(self, fake_funasr, fake_model_cache):
        w._get_spk_model()
        kwargs = fake_funasr.instances[0].kwargs
        assert Path(kwargs["model"]).is_absolute()
        assert Path(kwargs["model"]).is_dir()

    def test_worker_replies_models_missing_not_crash(self, monkeypatch, tmp_path):
        """断网且权重缺失 → worker 单请求回包 models_missing，常驻进程不崩。"""
        import io
        import json

        # main() 会以 build_engine_env() 重建 MODELSCOPE_CACHE → 一并桩掉
        monkeypatch.setattr(
            "core.engine_host.build_engine_env",
            lambda: {"MODELSCOPE_CACHE": str(tmp_path / "no-cache")},
        )
        req = json.dumps({"id": "r1", "cmd": "infer", "params": {"task": "asr", "audio_path": "/nonexistent/x.wav"}})
        # audio 不存在 → 先 bad request；再直接触发模型加载路径：
        wav = _write_wav(tmp_path / "a.wav")
        req2 = json.dumps({"id": "r2", "cmd": "infer", "params": {"task": "asr", "audio_path": wav}})
        req3 = json.dumps({"id": "r3", "cmd": "stop"})
        stdin = io.StringIO(req + "\n" + req2 + "\n" + req3 + "\n")
        stdout = io.StringIO()
        monkeypatch.setattr(sys, "stdin", stdin)
        monkeypatch.setattr(sys, "stdout", stdout)
        assert w.main() == 0
        replies = [json.loads(line) for line in stdout.getvalue().splitlines() if line.strip()]
        r2 = next(r for r in replies if r["id"] == "r2")
        assert r2["ok"] is False
        assert r2["error_code"] == "models_missing"
        assert "下载" in r2["error"]


# ============================================================
# AC-D3 — DEF-QA-R1-03：引擎预热（方案①）
# ============================================================


class TestWarmup:
    def test_handle_warmup_loads_both_singletons(self, fake_funasr, fake_model_cache, monkeypatch):
        result = w._handle_warmup()
        assert result["ok"] is True
        assert result["warmed"] == ["asr", "spk"]
        assert result["seconds"] >= 0
        assert w._ASR_MODEL is not None and w._SPK_MODEL is not None
        # 幂等：二次调用不再新建 AutoModel
        n = len(fake_funasr.instances)
        w._handle_warmup()
        assert len(fake_funasr.instances) == n

    def test_warmup_timeout_declares_frozen_import_margin(self):
        from core.engine_host import ENGINE_WARMUP_TIMEOUT

        # 留量声明（派工单 §2 DEF-03）：须覆盖冻结态冷 import ≈86s + 双模型加载 ≈32s + 余量
        assert ENGINE_WARMUP_TIMEOUT >= 86 + 32

    def test_start_kicks_warmup_after_healthy_ping(self, monkeypatch):
        """start() 健康检查通过后必须后台预热（DEF-03 方案①接线契约）。"""
        import time

        import core.engine_host as eh

        class _FakeStdin:
            def write(self, s): pass
            def flush(self): pass

        class _FakeStdout:
            def readline(self): return ""

        class _FakeProc:
            pid = 4242
            stdin = _FakeStdin()
            stdout = _FakeStdout()
            def poll(self): return None
            # stop() 收尾（含 atexit）走优雅 wait → terminate → kill 链，替身需补齐 Popen 面
            def wait(self, timeout=None): return 0
            def terminate(self): pass
            def kill(self): pass

        host = eh.EngineHost()
        seen_cmds = []

        monkeypatch.setattr(eh, "check_required_models", lambda: {"ready": True, "missing": [], "tampered": []})
        monkeypatch.setattr(eh.subprocess, "Popen", lambda *a, **k: _FakeProc())

        orig_request = eh.EngineHost.request

        def fake_request(self, cmd, params=None, timeout=None):
            seen_cmds.append(cmd)
            if cmd == "ping":
                from core.engine_worker import PROTOCOL_VERSION
                return {"ok": True, "id": "r", "protocol": PROTOCOL_VERSION}
            return {"ok": True, "id": "r", "seconds": 0.1, "warmed": ["asr", "spk"]}

        monkeypatch.setattr(eh.EngineHost, "request", fake_request)
        result = host.start()
        assert result["ok"] is True
        # warmup 在后台线程发送 → 等待其落地
        for _ in range(50):
            if host._warmup_state == "done":
                break
            time.sleep(0.02)
        assert "ping" in seen_cmds
        assert "warmup" in seen_cmds
        assert host._warmup_state == "done"
        monkeypatch.setattr(eh.EngineHost, "request", orig_request)


# ============================================================
# AC-D4 — DEF-QA-R1-04：retranscribe engine_override=cloud
# ============================================================


class _FakeBgTasks:
    """FastAPI BackgroundTasks 替身：记录 add_task 调用。"""

    def __init__(self):
        self.added = []

    def add_task(self, fn, *args, **kwargs):
        self.added.append((fn, args))


class TestEngineOverride:
    @pytest.fixture(autouse=True)
    def _no_disk_write(self, monkeypatch):
        """端点测试不触碰真实 data/tasks（save_task_to_disk / joblog 打桩）。"""
        import app.routers.tasks as tr
        import core.joblog
        import core.pipeline_runner as pr

        monkeypatch.setattr(tr, "save_task_to_disk", lambda tid: None)
        monkeypatch.setattr(pr, "save_task_to_disk", lambda tid: None)
        monkeypatch.setattr(core.joblog, "append_event", lambda *a, **k: None)

    def _make_task(self, tmp_path, tasks_dict, status="failed"):
        audio = _write_wav(tmp_path / "rec.wav")
        tid = "t-override-1"
        tasks_dict[tid] = {
            "task_id": tid, "status": status, "audio_path": audio, "audio_name": "rec.wav",
            "user_notes": "重要笔记：保留我", "dialogue": None, "summary": None,
        }
        return tid

    def test_pipeline_passes_override_to_transcribe(self, tmp_path, monkeypatch):
        import core.pipeline_runner as pr

        tid = self._make_task(tmp_path, pr.tasks)
        captured = {}

        def fake_stage1(audio_path, config=None, on_progress=None, speaker_count=None, engine_override=None):
            captured["engine_override"] = engine_override
            return {"dialogue": [{"speaker_id": 0, "text": "x", "sentences": []}], "transcription": {},
                    "audio_duration": 1.0, "normalized_path": str(audio_path)}

        monkeypatch.setattr(pr, "run_pipeline_stage1", fake_stage1)
        monkeypatch.setattr(pr, "_run_stage2_for_task", lambda *a, **k: captured.__setitem__("stage2", True))
        pr.run_pipeline_task(tid, Path(tmp_path / "rec.wav"), "rec.wav", engine_override="cloud")
        assert captured["engine_override"] == "cloud"

    def test_endpoint_forwards_cloud_and_keeps_notes(self, tmp_path, monkeypatch):
        import app.routers.tasks as tr

        bg = _FakeBgTasks()
        tid = self._make_task(tmp_path, tr.tasks)
        monkeypatch.setattr("core.cloud_client.is_cloud_enabled", lambda: True)
        resp = tr.retranscribe_task(tid, bg, tr.RetranscribeRequest(engine_override="cloud"))
        assert resp["status"] == "processing"
        fn, args = bg.added[0]
        assert args[0] == tid and args[-1] == "cloud"  # engine_override 末位透传
        # 笔记与音频状态不丢失（AC-4 Done 判据：不破坏「任务状态不丢失」）
        assert tr.tasks[tid]["user_notes"] == "重要笔记：保留我"
        assert tr.tasks[tid]["audio_path"]

    def test_endpoint_rejects_unknown_override(self, tmp_path, monkeypatch):
        from fastapi import HTTPException

        import app.routers.tasks as tr

        tid = self._make_task(tmp_path, tr.tasks)
        with pytest.raises(HTTPException) as ei:
            tr.retranscribe_task(tid, _FakeBgTasks(), tr.RetranscribeRequest(engine_override="local"))
        assert ei.value.status_code == 400

    def test_endpoint_rejects_cloud_when_not_configured(self, tmp_path, monkeypatch):
        from fastapi import HTTPException

        import app.routers.tasks as tr

        tid = self._make_task(tmp_path, tr.tasks)
        monkeypatch.setattr("core.cloud_client.is_cloud_enabled", lambda: False)
        with pytest.raises(HTTPException) as ei:
            tr.retranscribe_task(tid, _FakeBgTasks(), tr.RetranscribeRequest(engine_override="cloud"))
        assert ei.value.status_code == 400

    def test_endpoint_no_body_backward_compatible(self, tmp_path, monkeypatch):
        """无 body（旧前端）→ 按当前设置重跑，engine_override=None。"""
        import app.routers.tasks as tr

        tid = self._make_task(tmp_path, tr.tasks)
        bg = _FakeBgTasks()
        tr.retranscribe_task(tid, bg, None)
        assert bg.added[0][1][-1] is None


# ============================================================
# AC-D2 — DEF-QA-R1-02：docx 导出
# ============================================================


_SAMPLE_MD = """\
# 会议纪要

**会议时间**：2026-09-24 15:00
**参会人员**：2 位（Speaker 1/Speaker 2）

---

## 一、会议概要
本次围绕**预算编制**展开。

## 四、议题归类

### 议题1：架构设计
1. 第一步建立数据模型。
2. 分十步推进。

---

【Speaker 1】我们先讲架构。
【Speaker 2】好的。

*本纪要由 AI 自动生成，仅供参考。*
"""


class TestDocxExport:
    def test_markdown_to_docx_structure(self, tmp_path):
        from docx import Document

        from core.docx_export import markdown_to_docx

        out = markdown_to_docx(_SAMPLE_MD, tmp_path / "s.docx")
        assert Path(out).exists()
        doc = Document(out)
        styles = [p.style.name for p in doc.paragraphs]
        texts = [p.text for p in doc.paragraphs]
        joined = "\n".join(texts)
        # 标题映射（Heading 1/2/3）
        assert "Heading 1" in styles and "Heading 2" in styles and "Heading 3" in styles
        # 有序列表映射
        assert "List Number" in styles
        # 中文原文完整、说话人段落保留
        assert "【Speaker 1】我们先讲架构。" in joined
        assert "一、会议概要" in joined
        # 段内加粗 run
        bold_run = next(r for p in doc.paragraphs for r in p.runs if r.bold and r.text == "预算编制")
        assert bold_run.text == "预算编制"
        # 斜体脚注
        assert any(r.italic and "AI 自动生成" in r.text for p in doc.paragraphs for r in p.runs)
        # 文件为合法 zip（docx 容器 magic）
        assert Path(out).read_bytes()[:2] == b"PK"

    def test_cjk_eastasia_font_set(self, tmp_path):
        from docx import Document

        from core.docx_export import markdown_to_docx

        out = markdown_to_docx(_SAMPLE_MD, tmp_path / "f.docx")
        doc = Document(out)
        normal = doc.styles["Normal"]
        east = normal.element.get_or_add_rPr().get_or_add_rFonts().get(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia"
        )
        assert east == "宋体"

    def test_download_endpoint_format(self, tmp_path, monkeypatch):
        import app.routers.tasks as tr

        out_md = tmp_path / "minutes.md"
        out_md.write_text(_SAMPLE_MD, encoding="utf-8")
        tid = "t-docx-1"
        tr.tasks[tid] = {"task_id": tid, "status": "completed", "output_path": str(out_md)}

        md_resp = tr.download_summary(tid, format="md")
        assert md_resp.media_type == "text/markdown"

        docx_resp = tr.download_summary(tid, format="docx")
        assert "wordprocessingml" in docx_resp.media_type
        assert Path(docx_resp.path).exists() and Path(docx_resp.path).suffix == ".docx"
        assert Path(docx_resp.path).read_bytes()[:2] == b"PK"

        with pytest.raises(Exception) as ei:
            tr.download_summary(tid, format="pdf")
        from fastapi import HTTPException
        assert isinstance(ei.value, HTTPException) and ei.value.status_code == 400

        tr.tasks.pop(tid)
